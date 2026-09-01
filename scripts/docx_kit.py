# -*- coding: utf-8 -*-
"""python-docx 中文学术译文排版组件库（wmw-paper / Track A）

用法（在 <paper>_build.py 中）:
    import docx_kit as K
    from mypaper_content import META, TRANSLATOR_NOTE, BLOCKS, TABLES

    K.configure(fig_dir="my_extract", eq_dir="my_extract/eq", tables=TABLES,
                text_w=15.8, eq_dpi=360.0)
    doc = Document()
    K.setup(doc)
    K.title_block(doc, META, TRANSLATOR_NOTE)
    stat = K.render_blocks(doc, BLOCKS)
    out = K.save(doc, r"D:\\...\\文献目录", "xxx_中文翻译.docx")
    print("已生成:", out); print("统计:", stat)

块类型见 render_blocks；行内标记 *斜体* ^上标^ ~下标~。
"""
import sys
import os
import re

try:
    sys.stdout.reconfigure(encoding="utf-8")   # 幂等；库模块不可重新包装 sys.stdout
except Exception:
    pass
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

CN_FONT = "微软雅黑"
EN_FONT = "Times New Roman"
DARK_BLUE = RGBColor(0x1F, 0x3B, 0x73)
GRAY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

CFG = {
    "fig_dir": None,          # 插图目录
    "eq_dir": None,           # 公式图像目录
    "tables": {},             # {"T1": {"header": [...], "rows": [[...], ...]}}
    "text_w": 15.8,           # 版心宽 cm
    "eq_dpi": 360.0,          # 公式图像 DPI，须与裁剪时一致
    "col_widths": [4.6, 3.3, 4.6, 3.3],
    "body_size": 10.5,
    "body_line": 19,
}


def configure(**kw):
    CFG.update(kw)


def img_path(fname, eq=False):
    base = CFG["eq_dir"] if eq else CFG["fig_dir"]
    return os.path.join(base, fname)


def native_cm(fname):
    """公式图像按裁剪 DPI 还原原尺寸，预留编号位置"""
    w, _ = Image.open(img_path(fname, eq=True)).size
    return min(w / CFG["eq_dpi"] * 2.54, CFG["text_w"] - 1.6)


# ---------------- 底层：字体与段落装饰 ----------------

def set_run_font(r, size=10.5, bold=False, italic=False, color=None, cn=CN_FONT, en=EN_FONT):
    r.font.name = en
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color is not None:
        r.font.color.rgb = color
    rPr = r._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)


def shade_para(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def border_para(p, sides=("bottom",), color="1F3B73", sz="12"):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    for side in sides:
        el = OxmlElement("w:" + side)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), color)
        pBdr.append(el)
    pPr.append(pBdr)


def add_marked_runs(p, text, size, base_italic=False, color=None, bold=False):
    """解析 *斜体* ^上标^ ~下标~ 标记"""
    for tok in re.split(r"(\*[^*]+\*|\^[^^]+\^|~[^~]+~)", text):
        if not tok:
            continue
        if tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = p.add_run(tok[1:-1]); set_run_font(r, size, bold, True, color)
        elif tok.startswith("^") and tok.endswith("^") and len(tok) > 2:
            r = p.add_run(tok[1:-1]); set_run_font(r, size, bold, base_italic, color)
            r.font.superscript = True
        elif tok.startswith("~") and tok.endswith("~") and len(tok) > 2:
            r = p.add_run(tok[1:-1]); set_run_font(r, size, bold, base_italic, color)
            r.font.subscript = True
        else:
            r = p.add_run(tok); set_run_font(r, size, bold, base_italic, color)


# ---------------- 段落与标题 ----------------

def para(doc, text, size=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True,
         space_after=6, line=None, color=None, bold=False):
    size = CFG["body_size"] if size is None else size
    line = CFG["body_line"] if line is None else line
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line)
    if indent:
        pf.first_line_indent = Pt(size * 2)
    add_marked_runs(p, text, size, color=color, bold=bold)
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(14); pf.space_after = Pt(8)
    pf.keep_with_next = True
    add_marked_runs(p, text, 14, color=DARK_BLUE, bold=True)
    border_para(p, ("bottom",), "1F3B73", "8")
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(10); pf.space_after = Pt(6)
    pf.keep_with_next = True
    add_marked_runs(p, text, 12, color=DARK_BLUE, bold=True)
    return p


# ---------------- 图 / 公式 / 表 / 伪代码 ----------------

def formula_img(doc, num):
    """编号公式：居中制表位放图像 + 右制表位放 (num)"""
    fname = "eq" + num.replace(".", "_") + ".png"
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(5); pf.space_after = Pt(5)
    pf.tab_stops.add_tab_stop(Cm(CFG["text_w"] / 2), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(CFG["text_w"]), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    p.add_run().add_picture(img_path(fname, eq=True), width=Cm(native_cm(fname)))
    set_run_font(p.add_run("\t(" + num + ")"), 10.5, False, False, GRAY)
    return p


def figure(doc, spec):
    """spec = "文件名|图注|宽度cm" """
    fname, caption, width_cm = spec.split("|")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True          # 图与图注不得分页
    p.add_run().add_picture(img_path(fname), width=Cm(float(width_cm)))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    cap.paragraph_format.line_spacing = Pt(14)
    add_marked_runs(cap, caption, 9, color=GRAY)


def build_table(doc, spec):
    """spec = ("表 1　题名", "T1")；整表不跨页"""
    title, key = spec
    data = CFG["tables"][key]
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(10)
    cap.paragraph_format.space_after = Pt(3)
    cap.paragraph_format.keep_with_next = True
    add_marked_runs(cap, title, 10.5, color=DARK_BLUE, bold=True)

    ncol = len(data["header"])
    tab = doc.add_table(rows=1, cols=ncol)
    tab.style = "Table Grid"
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tab.rows[0].cells
    for c, cell in enumerate(hdr):
        cell.text = ""
        pp = cell.paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pp.paragraph_format.space_after = Pt(1)
        add_marked_runs(pp, data["header"][c], 9.5, color=WHITE, bold=True)
        shade_cell(cell, "1F3B73")
    for row in data["rows"]:
        cells = tab.add_row().cells
        for c, cell in enumerate(cells):
            cell.text = ""
            pp = cell.paragraphs[0]
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pp.paragraph_format.space_after = Pt(1)
            add_marked_runs(pp, row[c], 9.5)
    widths = CFG["col_widths"]
    for row in tab.rows:
        for c, cell in enumerate(row.cells):
            cell.width = Cm(widths[c])
    # 防跨页：行内不拆分 + 除末行外均与下一段同页
    for row in tab.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for row in tab.rows[:-1]:
        for cell in row.cells:
            for pp in cell.paragraphs:
                pp.paragraph_format.keep_with_next = True
    q = doc.add_paragraph()
    q.paragraph_format.space_after = Pt(4)            # 表后必须留段，否则与后续内容粘连
    return tab


def algo(doc, lines):
    """伪代码：16pt 精确行距，行号与全角缩进原样保留"""
    for ln in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(1)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(16)
        add_marked_runs(p, ln, 10)
    e = doc.add_paragraph()
    e.paragraph_format.space_after = Pt(8)


# ---------------- 页面与页脚 ----------------

def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    for el in (fld1, instr, fld2):
        r = p.add_run(); set_run_font(r, 9, False, False, GRAY)
        r._element.append(el)


def setup(doc, margin=(2.5, 2.3, 2.6, 2.6)):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(margin[0]); sec.bottom_margin = Cm(margin[1])
    sec.left_margin = Cm(margin[2]); sec.right_margin = Cm(margin[3])
    add_page_number_footer(sec)
    return sec


# ---------------- 标题区 ----------------

def title_block(doc, meta, translator_note=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run(meta.get("field", "")), 10, True, False, GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_marked_runs(p, meta["title_cn"], 17, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_marked_runs(p, meta.get("title_en", ""), 10.5, color=GRAY)
    for r in p.runs:
        r.font.italic = True

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_marked_runs(p, meta.get("authors", ""), 10.5)

    for aff in meta.get("affiliations", []):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
        add_marked_runs(p, aff, 9, color=GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(1)
    add_marked_runs(p, meta.get("journal", ""), 9.5, color=GRAY)
    if meta.get("doi"):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        add_marked_runs(p, meta["doi"], 9.5, color=GRAY)

    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
    add_marked_runs(p, meta.get("note", ""), 10, color=DARK_BLUE, bold=True)
    border_para(p, ("top", "bottom"), "1F3B73", "6")

    if translator_note:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(16)
        add_marked_runs(p, translator_note, 9.5, color=GRAY)
        shade_para(p, "FBF6E9")


# ---------------- 块渲染 ----------------

def render_blocks(doc, blocks):
    stat = {}
    for block in blocks:
        kind = block[0]
        if kind == "eqimg":
            formula_img(doc, block[1])
        elif kind == "fig":
            figure(doc, block[1])
        elif kind == "table":
            build_table(doc, block[1])
        elif kind == "algo":
            algo(doc, block[1])
        elif kind == "h1":
            h1(doc, block[1])
        elif kind == "h2":
            h2(doc, block[1])
        elif kind == "li":
            para(doc, block[1], indent=False, space_after=4)
        elif kind == "note":
            para(doc, block[1], size=9.5, indent=False, color=GRAY, space_after=8, line=15)
        elif kind == "kw":
            p = para(doc, block[1], size=10.5, indent=False, space_after=10, line=18)
            for r in p.runs:
                r.font.bold = True
        elif kind == "abs":
            p = para(doc, block[1], size=10.5, space_after=8, line=18)
            shade_para(p, "F2F2F2")
            border_para(p, ("left",), "1F3B73", "18")
        else:
            para(doc, block[1])
        stat[kind] = stat.get(kind, 0) + 1
    return stat


# ---------------- 保存 ----------------

def save(doc, out_dir, fname, fallback=None):
    """优先写目标目录（如源 PDF 同目录），被占用则回退"""
    out = os.path.join(out_dir, fname)
    try:
        doc.save(out)
    except OSError:
        out = os.path.join(fallback or os.getcwd(), fname)
        doc.save(out)
    return out
